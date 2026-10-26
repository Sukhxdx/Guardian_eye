"""Build the GuardianEye Multimodal AI (MAI) project PowerPoint.

Focuses on multimodal architecture and late fusion. Does not invent detection metrics.
"""

from __future__ import annotations

import json
import shutil
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SUMMARY_PATH = ROOT / "reports" / "bda" / "summary.json"
OUT_DOCS = ROOT / "docs" / "MAI_Project_Presentation_2026_2027.pptx"
OUT_ROOT = ROOT / "MAI_Project_Presentation_2026_2027.pptx"

NAVY = RGBColor(0x0B, 0x1D, 0x36)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED_GREEN = RGBColor(0x3D, 0x7A, 0x5F)
SLATE = RGBColor(0x4A, 0x55, 0x68)
TABLE_ALT = RGBColor(0xE8, 0xEE, 0xF4)

MEMBERS = [
    "Akhil Chivakula — 23070126009",
    "Dipti Kothari — 23070126040",
    "Khurinji Malar — 23070126059",
    "Sukhada Dhingra — 24070126505",
]

GUIDE = "Dr. Nivedita Mishra"
HOD = "Dr. Sumit Kumar"
TOTAL = 13


def _set_run(run, text, size=18, bold=False, color=NAVY, font="Calibri"):
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=NAVY, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    _set_run(p.add_run(), text, size=size, bold=bold, color=color)
    return box


def _solid_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _add_rect(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    _solid_fill(shape, color)
    return shape


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _header(slide, title, subtitle=None):
    _add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(0.95), NAVY)
    _add_rect(slide, Inches(0), Inches(0.95), Inches(13.333), Inches(0.06), MUTED_GREEN)
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.22), Inches(12.3), Inches(0.55))
    _set_run(box.text_frame.paragraphs[0].add_run(), title, size=26, bold=True, color=WHITE)
    if subtitle:
        sub = slide.shapes.add_textbox(Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.4))
        _set_run(sub.text_frame.paragraphs[0].add_run(), subtitle, size=13, color=SLATE)


def _footer(slide, page):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(7.15), Inches(10), Inches(0.3))
    _set_run(
        box.text_frame.paragraphs[0].add_run(),
        "GuardianEye | Multimodal AI (MAI) | Group 27 | SIT AIML | AY 2026–2027",
        size=10,
        color=SLATE,
    )
    num = slide.shapes.add_textbox(Inches(11.5), Inches(7.15), Inches(1.3), Inches(0.3))
    num.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    _set_run(num.text_frame.paragraphs[0].add_run(), f"{page} / {TOTAL}", size=10, color=SLATE)


def _bullets(slide, left, top, width, height, items, size=16):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        _set_run(p.add_run(), f"•  {item}", size=size, color=NAVY)
    return box


def _fill_cell(cell, text, size=12, bold=False, color=NAVY, fill=None):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fill if fill is not None else WHITE
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.clear()
    _set_run(p.add_run(), text, size=size, bold=bold, color=color)


def _table(slide, left, top, width, height, headers, rows):
    table = slide.shapes.add_table(1 + len(rows), len(headers), left, top, width, height).table
    for i, h in enumerate(headers):
        _fill_cell(table.cell(0, i), h, size=11, bold=True, color=WHITE, fill=NAVY)
    for r, row in enumerate(rows):
        bg = TABLE_ALT if r % 2 else WHITE
        for c, val in enumerate(row):
            _fill_cell(table.cell(r + 1, c), str(val), size=11, fill=bg)
    return table


def _load():
    if SUMMARY_PATH.is_file():
        return json.loads(SUMMARY_PATH.read_text())
    return {}


def build() -> Path:
    s = _load()
    img = s.get("image", {})
    aud = s.get("audio", {})
    splits = s.get("audio_splits", {}).get("by_experiment_and_split", {})
    proxy = s.get("synthetic_proxy_pairs", {})
    levels = proxy.get("by_expected_level", {"HIGH": 6, "MEDIUM": 6, "LOW": 4})

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # 1 Title
    slide = _blank(prs)
    _add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
    _add_rect(slide, Inches(0), Inches(0), Inches(0.18), Inches(7.5), MUTED_GREEN)
    _add_textbox(
        slide,
        Inches(0.7),
        Inches(1.2),
        Inches(12),
        Inches(1.0),
        "GuardianEye",
        size=40,
        bold=True,
        color=WHITE,
    )
    _add_textbox(
        slide,
        Inches(0.7),
        Inches(2.1),
        Inches(12),
        Inches(0.8),
        "Multimodal Late Fusion of Image and Audio Poaching Indicators",
        size=22,
        color=WHITE,
    )
    _add_textbox(
        slide,
        Inches(0.7),
        Inches(3.1),
        Inches(12),
        Inches(0.4),
        "Multimodal AI (MAI)  ·  Group 27  ·  Academic Year 2026–2027",
        size=16,
        color=MUTED_GREEN,
    )
    _add_textbox(
        slide,
        Inches(0.7),
        Inches(3.7),
        Inches(12),
        Inches(0.35),
        f"Guide: {GUIDE}, Assistant Professor, AIML",
        size=14,
        color=WHITE,
    )
    _add_textbox(
        slide,
        Inches(0.7),
        Inches(4.1),
        Inches(12),
        Inches(0.35),
        f"Head of the Department: {HOD}",
        size=14,
        color=WHITE,
    )
    _add_textbox(
        slide,
        Inches(0.7),
        Inches(4.7),
        Inches(12),
        Inches(1.4),
        "\n".join(MEMBERS),
        size=14,
        color=WHITE,
    )
    _add_textbox(
        slide,
        Inches(0.7),
        Inches(6.5),
        Inches(12),
        Inches(0.35),
        "Symbiosis Institute of Technology  ·  B.Tech AIML",
        size=13,
        color=SLATE,
    )

    # 2 Problem
    slide = _blank(prs)
    _header(slide, "Problem Formulation", "Why multimodal AI for wildlife indicators")
    _bullets(
        slide,
        Inches(0.6),
        Inches(1.4),
        Inches(12),
        Inches(5.2),
        [
            "Patrol evidence arrives as RGB stills, aerial thermal frames, and short audio clips.",
            "Streams are rarely time-aligned; a single sensor reading is incomplete.",
            "No labelled poacher class in the configured collections.",
            "One detection is an indicator, not proof of poaching.",
            "Video is sampled into frames and follows the image branch — not a third modality.",
            "Late fusion combines unimodal scores at decision time, not in a joint encoder.",
        ],
        size=17,
    )
    _footer(slide, 2)

    # 3 Objectives
    slide = _blank(prs)
    _header(slide, "Objectives", "Completed multimodal work only")
    _bullets(
        slide,
        Inches(0.6),
        Inches(1.4),
        Inches(12),
        Inches(5.2),
        [
            "Train separate visual detectors for aerial_thermal and camera_trap_rgb (YOLOv8).",
            "Classify gunshot vs background on log-mel windows with a convolutional network.",
            "Evaluate audio under Belize↔Vietnam transfer and in-domain controls.",
            "Fuse modalities with rules, time windows, and missing-modality degradation.",
            "Build synthetic proxy pairs when field-aligned pairs do not exist.",
            "Serve a multimodal API and Gradio dashboard with an honest audit log.",
        ],
        size=17,
    )
    _footer(slide, 3)

    # 4 Architecture
    slide = _blank(prs)
    _header(slide, "Multimodal Architecture", "Two unimodal branches + late fusion")
    # three cards
    cards = [
        (0.5, "Image branch", ["YOLOv8 per domain", "human / vehicle / animal / empty", "Empty frames retained"]),
        (4.7, "Audio branch", ["16 kHz · 3 s windows", "64-bin log-mel CNN", "Class-weighted loss"]),
        (8.9, "Late fusion", ["Rules mode (primary)", "5 s coincidence window", "HIGH / MEDIUM / LOW"]),
    ]
    for left, title, items in cards:
        _add_rect(slide, Inches(left), Inches(1.5), Inches(3.8), Inches(4.6), TABLE_ALT)
        _add_rect(slide, Inches(left), Inches(1.5), Inches(3.8), Inches(0.55), MUTED_GREEN)
        _add_textbox(slide, Inches(left + 0.2), Inches(1.58), Inches(3.4), Inches(0.4), title, size=16, bold=True, color=WHITE)
        _bullets(slide, Inches(left + 0.2), Inches(2.3), Inches(3.4), Inches(3.5), items, size=14)
    _footer(slide, 4)

    # 5 Visual modality
    slide = _blank(prs)
    _header(slide, "Visual Modality", "Domain-separated object detection")
    _bullets(
        slide,
        Inches(0.6),
        Inches(1.4),
        Inches(12),
        Inches(2.2),
        [
            "Independent YOLOv8 runs for aerial_thermal and camera_trap_rgb.",
            "Thermal uses channel replicate into three identical channels for an RGB backbone.",
            "Rare classes oversampled; empty negatives kept.",
            "Visual indicators for fusion: human and vehicle only.",
        ],
        size=16,
    )
    by_cls = img.get("by_domain_and_class", {})
    at = by_cls.get("aerial_thermal", {"human": 2, "vehicle": 1, "animal": 2, "empty": 6})
    ct = by_cls.get("camera_trap_rgb", {"human": 2, "vehicle": 1, "animal": 2, "empty": 6})
    _table(
        slide,
        Inches(0.6),
        Inches(3.8),
        Inches(12),
        Inches(2.2),
        ["Domain", "Human", "Vehicle", "Animal", "Empty", "Total"],
        [
            ["aerial_thermal", at.get("human", 2), at.get("vehicle", 1), at.get("animal", 2), at.get("empty", 6), img.get("by_domain", {}).get("aerial_thermal", 11)],
            ["camera_trap_rgb", ct.get("human", 2), ct.get("vehicle", 1), ct.get("animal", 2), ct.get("empty", 6), img.get("by_domain", {}).get("camera_trap_rgb", 11)],
        ],
    )
    _footer(slide, 5)

    # 6 Audio modality
    slide = _blank(prs)
    _header(slide, "Audio Modality", "Site-stratified acoustic classification")
    site = aud.get("by_site_and_label", {"belize": {"gunshot": 4, "background": 10}, "vietnam": {"gunshot": 5, "background": 10}})
    _bullets(
        slide,
        Inches(0.6),
        Inches(1.4),
        Inches(12),
        Inches(2.0),
        [
            "Log-mel CNN classifies gunshot vs background.",
            "Background and distractor impulses retained.",
            "Cross-site experiments measure geographic transfer into fusion.",
        ],
        size=16,
    )
    _table(
        slide,
        Inches(0.6),
        Inches(3.6),
        Inches(12),
        Inches(2.4),
        ["Site", "Gunshot", "Background", "Total"],
        [
            ["Belize", site["belize"]["gunshot"], site["belize"]["background"], aud.get("by_site", {}).get("belize", 14)],
            ["Vietnam", site["vietnam"]["gunshot"], site["vietnam"]["background"], aud.get("by_site", {}).get("vietnam", 15)],
        ],
    )
    _footer(slide, 6)

    # 7 Cross-site
    slide = _blank(prs)
    _header(slide, "Cross-Site Audio Experiments", "Transfer both ways + in-domain controls")
    rows = []
    for name in ("belize_to_vietnam", "vietnam_to_belize", "in_domain_belize", "in_domain_vietnam"):
        vals = splits.get(name, {"train": "—", "val": "—", "test": "—"})
        rows.append([name, vals.get("train"), vals.get("val"), vals.get("test")])
    _table(
        slide,
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(3.5),
        ["Experiment", "Train", "Validation", "Test"],
        rows,
    )
    _add_textbox(
        slide,
        Inches(0.6),
        Inches(5.3),
        Inches(12),
        Inches(1.2),
        "Validation stays on the training site. Cross-site tests use the other geography entirely. "
        "Fusion inherits whatever site shift the audio encoder exhibits.",
        size=15,
        color=SLATE,
    )
    _footer(slide, 7)

    # 8 Late fusion
    slide = _blank(prs)
    _header(slide, "Late Fusion Policy", "configs/fusion.yaml · rules mode")
    _bullets(
        slide,
        Inches(0.6),
        Inches(1.4),
        Inches(12),
        Inches(5.2),
        [
            "HIGH — human/vehicle and gunshot, both ≥ 0.50, inside a 5 s window → default alert.",
            "MEDIUM — exactly one strong indicator, or both outside the window.",
            "LOW — neither indicator.",
            "Missing modality: unimodal thresholds 0.85; level cannot be HIGH.",
            "Weight profiles (balanced / night / wind) available for weighted scoring.",
            "Animal detections are logged and excluded from poaching indicators.",
        ],
        size=16,
    )
    _footer(slide, 8)

    # 9 Synthetic proxy
    slide = _blank(prs)
    _header(slide, "Synthetic Proxy Multimodal Pairs", "Not field-aligned incidents")
    _bullets(
        slide,
        Inches(0.6),
        Inches(1.4),
        Inches(6.0),
        Inches(4.5),
        [
            f"{proxy.get('records', 16)} constructed pairs.",
            "Protocol tag: synthetic_proxy on every row.",
            "Eight scenarios × two visual domains.",
            "Used to compare image-only, audio-only, and fused decisions.",
            "Explicitly labelled as a proxy protocol in reports.",
        ],
        size=16,
    )
    _table(
        slide,
        Inches(7.0),
        Inches(1.5),
        Inches(5.5),
        Inches(3.2),
        ["Expected level", "Pairs"],
        [["HIGH", levels.get("HIGH", 6)], ["MEDIUM", levels.get("MEDIUM", 6)], ["LOW", levels.get("LOW", 4)]],
    )
    _footer(slide, 9)

    # 10 Serving
    slide = _blank(prs)
    _header(slide, "Multimodal Serving", "FastAPI + Gradio + SQLite audit")
    _bullets(
        slide,
        Inches(0.6),
        Inches(1.4),
        Inches(12),
        Inches(5.2),
        [
            "API on port 8742; dashboard at /dashboard accepts image and audio together.",
            "GET /health — process liveness.",
            "GET /model-info — per-branch checkpoint honesty (loaded vs none).",
            "Every decision logged to data/events.sqlite, including non-alerts.",
            "Banner: proof-of-concept. Indicators, not confirmed poaching.",
            "Current tree: no checkpoints in models/; evidence marked unreliable.",
        ],
        size=16,
    )
    _footer(slide, 10)

    # 11 Corpus snapshot
    slide = _blank(prs)
    _header(slide, "Multimodal Corpus Snapshot", "Metadata counts only — not detection scores")
    _table(
        slide,
        Inches(0.6),
        Inches(1.5),
        Inches(12),
        Inches(4.5),
        ["Quantity", "Value"],
        [
            ["Image records", img.get("records", 22)],
            ["Empty images retained", img.get("empty_images_retained", 12)],
            ["Audio windows", aud.get("records", 29)],
            ["Background clips retained", aud.get("background_clips_retained", 20)],
            ["Audio split assignment rows", s.get("audio_splits", {}).get("records", 84)],
            ["Synthetic proxy pairs", proxy.get("records", 16)],
            ["Checkpoints in models/", "none"],
        ],
    )
    _footer(slide, 11)

    # 12 Conclusion
    slide = _blank(prs)
    _header(slide, "Conclusion", "Multimodal AI deliverable")
    _bullets(
        slide,
        Inches(0.6),
        Inches(1.4),
        Inches(12),
        Inches(5.2),
        [
            "GuardianEye implements late multimodal fusion of vision and audio indicators.",
            "Unimodal branches stay specialised; fusion is rules-based and time-aware.",
            "Site stratification and empty/background retention keep evaluation honest.",
            "Synthetic proxy pairs stand in where field-aligned multimodal data do not exist.",
            "Serving exposes multimodal uploads with explicit reliability reporting.",
            "No fabricated mAP, F1, precision, or recall — metrics require checkpoints.",
        ],
        size=16,
    )
    _footer(slide, 12)

    # 13 Thank you
    slide = _blank(prs)
    _add_rect(slide, Inches(0), Inches(0), Inches(13.333), Inches(7.5), NAVY)
    _add_rect(slide, Inches(0), Inches(0), Inches(0.18), Inches(7.5), MUTED_GREEN)
    _add_textbox(slide, Inches(0.7), Inches(2.0), Inches(12), Inches(0.7), "Thank you", size=40, bold=True, color=WHITE)
    _add_textbox(slide, Inches(0.7), Inches(2.9), Inches(12), Inches(0.4), "Questions welcome", size=20, color=MUTED_GREEN)
    _add_textbox(slide, Inches(0.7), Inches(3.7), Inches(12), Inches(1.5), "\n".join(MEMBERS), size=15, color=WHITE)
    _add_textbox(
        slide,
        Inches(0.7),
        Inches(5.5),
        Inches(12),
        Inches(0.8),
        f"Guide: {GUIDE}  ·  Head of the Department: {HOD}\nMultimodal AI (MAI)  ·  Group 27  ·  SIT AIML  ·  AY 2026–2027",
        size=14,
        color=WHITE,
    )

    OUT_DOCS.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_DOCS)
    shutil.copy2(OUT_DOCS, OUT_ROOT)
    return OUT_DOCS


if __name__ == "__main__":
    path = build()
    print(f"wrote {path}")
    print(f"wrote {OUT_ROOT}")
